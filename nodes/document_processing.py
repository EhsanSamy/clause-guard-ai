"""
Person 1 — node: document_processing
Owns: document_text

Reads case.raw_input (already confirmed valid by validate_input — either
a file path or pasted text), extracts/cleans plain text, and writes
document_text.

Per docs/guardrails.md §6 (Tool Failures): failures are appended to
tool_errors instead of crashing the graph. tool_errors uses an additive
reducer (agent/state.py: Annotated[List[str], operator.add]), so this
node returns only the NEW error(s) for this run, not the accumulated list.

Note on clause_id (docs/schema.md §2): that field is listed as
"Set by: document_processing", but a Clause object can't be created here
— clause_text/clause_tag are required and only known once extract_clauses
(Person 2) has actually split + tagged the text. So this module owns the
*ID format*, exposed as generate_clause_id() below, and extract_clauses
is expected to call it when constructing each Clause so IDs stay
consistent with how document_processing numbers things.
"""
from __future__ import annotations
import io
import os
import re
from agent.config import GUARDRAILS
from agent.state import GraphState, to_case_state, validate_update


def generate_clause_id(index: int) -> str:
    """Stable clause id generator, e.g. generate_clause_id(3) -> 'clause_003'.

    Owned by document_processing per docs/schema.md. Call this from
    extract_clauses (Person 2) once per clause, in extraction order,
    starting at index=1.
    """
    return f"clause_{index:03d}"


def _extract_pdf(raw_bytes: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw_bytes))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as e:
            raise ValueError("PDF is encrypted/password-protected") from e

    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    print(f"[DEBUG] Extracted text length: {len(text)} chars")
    print(f"[DEBUG] First 300 chars: {text[:300]!r}")
    if not text.strip():
        raise ValueError(
            "No extractable text found — looks like a scanned image PDF "
            "(would need OCR, which this node does not perform)."
        )
    return text


def _extract_docx(raw_bytes: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(raw_bytes))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\x0c", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def document_processing_node(state: GraphState) -> GraphState:
    case = to_case_state(state)
    if not case.input_valid:
        return {}  # validate_input already rejected this case — nothing to do

    raw = case.raw_input

    try:
        if os.path.exists(raw):
            ext = os.path.splitext(raw)[1].lower()
            with open(raw, "rb") as f:
                raw_bytes = f.read()

            if ext == ".pdf":
                text = _extract_pdf(raw_bytes)
            elif ext == ".docx":
                text = _extract_docx(raw_bytes)
            elif ext == ".txt":
                text = raw_bytes.decode("utf-8", errors="replace")
            else:
                raise ValueError(f"Unsupported extension: '{ext}'")
        else:
            text = raw  # pasted text — length already checked by validate_input

        clean = _clean_text(text)

        if len(clean) < GUARDRAILS.min_document_chars:
            raise ValueError(
                f"Extracted text too short ({len(clean)} chars, min "
                f"{GUARDRAILS.min_document_chars}) — likely a bad extraction."
            )
        if len(clean) > GUARDRAILS.max_document_chars:
            raise ValueError(
                f"Extracted text too large ({len(clean)} chars, max "
                f"{GUARDRAILS.max_document_chars})."
            )

        return validate_update({"document_text": clean})

    except Exception as e:
        return {"tool_errors": [f"document_processing: {e}"]}
